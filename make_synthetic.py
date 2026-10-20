#!/usr/bin/env python3
"""
Generate a fully FICTIONAL collegiate-taiko interview corpus for the RAG pipeline.

Everything here is invented: groups G101-G108, people P101+, schools S101+, their
histories, schedules and answers. The corpus mirrors the format of the private
research corpus the pipeline was built for, so every script runs end to end:

  data/synthetic/docs/*.docx      one email interview per group (header + Q:/A: pairs)
  data/synthetic/groups.csv       structured table of the same facts (table route)
  data/synthetic/eval_questions.csv   ~55 questions with gold answers + gold chunk IDs
  data/synthetic/eval_split.csv   stratified dev/test split

All facts are drawn once from a seeded random generator; the documents, the table
and the gold answers are all written from those same facts, so gold answers are
correct by construction. The documents also include some realistic messiness the
ingestion step has to handle: a roster table, a numbered list, a question wrapped
onto a second line, a two-part question answered only halfway, and questions a
group simply never answered ("not stated").

Usage:
    pip install python-docx
    python make_synthetic.py                 # writes data/synthetic/
    python make_synthetic.py --seed 7        # a different fictional world
"""

from __future__ import annotations

import argparse
import csv
import random
from pathlib import Path

try:
    import docx
except ImportError:
    raise SystemExit("python-docx is required: pip install python-docx")

N_GROUPS = 8
REGIONS = ["Northeast", "Midwest", "South", "West"]
MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August",
          "September", "October", "November", "December"]
DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]
TRADITIONAL = ["Renshuu", "Matsuri", "Hachijo", "Miyake", "Yatai Bayashi"]
ORIGINALS = ["River Song", "Lantern Walk", "First Snow", "Iron Bridge", "Night Market", "Ember"]
CONFERENCES = ["ITI", "ECTC", "NATC"]
PRACTICE_PLACES = ["the basement of the music building", "a dance studio in the recreation center",
                   "a multipurpose room in the student union", "an outdoor courtyard on campus"]
STORAGE = [("a storage closet in the music building", False), ("the advisor's office", True),
           ("a member's apartment", True), ("a shared equipment room in the student union", False),
           ("a garage owned by a former member", True)]
EQUIPMENT = {
    "built": "We built our own drums from wine barrels over two summers.",
    "purchased": "We bought our drums from a taiko maker using a student activities grant.",
    "donated": "A local community taiko group donated its older drums to us.",
    "makeshift": "We still don't have real taiko drums, so we practice on tires and buckets.",
}
# (challenge, adaptation, short phrase used in a "why" question)
CHALLENGES = [
    ("we lost our regular rehearsal room when the music building was renovated",
     "we moved practices to an outdoor courtyard", "move practices outdoors"),
    ("members had heavy lab and work schedules",
     "we switched to fewer but longer weekend sessions", "switch to fewer, longer sessions"),
    ("our drums were damaged when the storage room flooded",
     "we spent a year building replacement drums", "spend a year building drums"),
    ("several senior members graduated in the same year",
     "we started a mentorship program pairing new and returning members", "start a mentorship program"),
    ("a neighboring dorm complained about the noise",
     "we began practicing on tire drums and only use real drums right before shows",
     "start practicing on tire drums"),
    ("our faculty advisor left the university",
     "we reorganized as a fully student-run club with elected officers", "reorganize with elected officers"),
    ("we had no budget for travel",
     "we focused on campus performances instead of conferences", "focus on campus performances"),
    ("interest dropped during a year of remote classes",
     "we recorded video lessons so new members could learn at home", "record video lessons"),
]


class World:
    def __init__(self, seed: int):
        self.rng = random.Random(seed)
        self.next_person = 101
        self.groups = [self._make_group(i) for i in range(N_GROUPS)]
        self._add_links()

    def person(self) -> str:
        pid = f"P{self.next_person}"
        self.next_person += 1
        return pid

    def _make_group(self, i: int) -> dict:
        r = self.rng
        gid, school = f"G{101 + i}", f"S{101 + i}"
        interviewee = self.person()
        faculty_founded = i in (2, 6)                    # guarantee two faculty-founded groups
        founders = [interviewee] if faculty_founded else \
            ([interviewee] if r.random() < 0.4 else []) + [self.person() for _ in range(r.randint(1, 3))]
        year = 1996 + i * 3 + r.randint(0, 2)
        practice = ["once", "same_day_two", "twice", "three", "once", "twice", "once", "same_day_two"][i]
        g = {
            "group": gid, "school": school, "interviewee": interviewee,
            "region": REGIONS[(i + r.randint(0, 3)) % 4],
            "date": f"20{19 + (i % 2)}-{r.randint(1, 5) if i % 2 else r.randint(10, 12):02d}-{r.randint(1, 28):02d}",
            "founded_year": year, "founded_month": r.choice(MONTHS),
            "founder_type": "faculty" if faculty_founded else "students", "founders": founders,
            "run_by": "instructor" if faculty_founded else "student",
            "practice": practice, "practice_day": r.choice(DAYS),
            "practice_days2": r.sample(DAYS, 3),
            "practice_place": r.choice(PRACTICE_PLACES),
            "storage": STORAGE[i % len(STORAGE)],
            "auditions": ["none", "workshops", "none", "case_by_case", "workshops", "none", "none", "workshops"][i],
            "advisor": interviewee if faculty_founded else (self.person() if r.random() < 0.5 else None),
            "funding": r.choice(["yes", "no"]),
            "pieces": r.sample(TRADITIONAL, r.randint(1, 3)) + ([r.choice(ORIGINALS)] if r.random() < 0.4 else []),
            "membership_mode": ["per_year", "estimate", "roster", "per_year", "per_year", "estimate", "per_year",
                                "per_year"][i],
            "conferences": r.sample(CONFERENCES, r.randint(0, 2)),
            "challenge": CHALLENGES[i],
            "equipment": ["purchased", "built", "donated", "makeshift", "purchased", "built", "donated",
                          "purchased"][i],
            "half_answer_storage": i == 4,                   # two-part question, only storage answered
            "omit": set(),
            "collaborated_with": None, "founders_from": None,
            "founders_with_experience": [],
        }
        if "Renshuu" not in g["pieces"] and i in (0, 3, 5):
            g["pieces"].insert(0, "Renshuu")                 # make the Renshuu aggregate non-trivial
        years = list(range(max(year, 2016), 2020))
        base = r.randint(5, 9)
        g["membership"] = {y: base + k * r.randint(2, 6) for k, y in enumerate(years)}
        g["roster"] = [(self.person(), r.randint(year, 2019)) for _ in range(r.randint(6, 9))] \
            if g["membership_mode"] == "roster" else []
        g["estimate_total"] = r.randint(15, 40)
        # questions this group never answered -> "not stated"
        pool = ["advisor", "funding", "conferences"]
        g["omit"] = set(r.sample(pool, 1 if i % 2 else 2)) if i not in (2, 6) else {"conferences"}
        if g["advisor"] == interviewee and "advisor" in g["omit"]:
            g["omit"].discard("advisor")
        if i in (5, 7):                                  # guarantee groups that answered "no advisor"
            g["advisor"] = None
            g["omit"].discard("advisor")
        if i in (0, 3):                                  # guarantee groups that attended ITI
            g["conferences"] = sorted(set(g["conferences"]) | {"ITI"})
            g["omit"].discard("conferences")
        if not g["omit"]:
            g["omit"] = {"funding"}
        return g

    def _add_links(self):
        r = self.rng
        by_year = sorted(self.groups, key=lambda g: g["founded_year"])
        # founders of two newer groups came from an older group
        for newer in (by_year[5], by_year[7]):
            if newer["founder_type"] != "students":
                continue
            older = r.choice([g for g in by_year if g["founded_year"] < newer["founded_year"]])
            newer["founders_from"] = older["group"]
            newer["founders_with_experience"] = newer["founders"][:max(1, len(newer["founders"]) - 1)]
        # one-directional collaborations (only the first group mentions it)
        for a, b in [(by_year[3], by_year[1]), (by_year[6], by_year[2])]:
            a["collaborated_with"] = b["group"]


# ---------------------------------------------------------------------------
# Document writing
# ---------------------------------------------------------------------------

def schedule_text(g: dict) -> str:
    d = g["practice_day"]
    d1, d2, d3 = g["practice_days2"]
    return {
        "once": f"once a week on {d}s from 6-8 pm",
        "same_day_two": f"every {d}, with a beginner session from 6-7:30 pm and an advanced session from 7:30-9 pm",
        "twice": f"twice a week, on {d1}s and {d2}s, for two hours each",
        "three": f"three times a week ({d1}, {d2} and {d3}), about two hours each",
    }[g["practice"]]


def times_per_week(g: dict) -> int:
    return {"once": 1, "same_day_two": 1, "twice": 2, "three": 3}[g["practice"]]


def founders_text(g: dict) -> str:
    who = ", ".join(g["founders"])
    if g["founder_type"] == "faculty":
        return f"I founded {g['group']} when I joined the faculty at {g['school']}."
    count = ["zero", "one", "two", "three", "four", "five"][len(g["founders"])]
    s = f"{g['group']} was founded by {count} {g['school']} student" \
        f"{'s' if len(g['founders']) > 1 else ''}: {who}."
    if g["founders_from"]:
        exp = ", ".join(g["founders_with_experience"])
        s += f" {exp} had played with {g['founders_from']} before starting the group."
    else:
        s += " None of us had played taiko before."
    return s


def build_qas(g: dict, rng: random.Random) -> list[tuple[str, str, list]]:
    """Return [(key, question, answer_blocks)] in document order. An answer block is a
    paragraph string or ("table", rows)."""
    G = g["group"]
    qas = [("founders", f"who are the founder(s) of {G}?", [founders_text(g)]),
           ("founded", "when was the group founded?",
            [rng.choice([f"{G} was founded in {g['founded_month']} {g['founded_year']}.",
                         f"We started in {g['founded_month']} {g['founded_year']}."])]),
           ("run_by", "is the group student-run or instructor-run?",
            ["It is primarily instructor-run; I lead most of the practices." if g["run_by"] == "instructor"
             else "It is completely student-run. Officers are elected every spring."]),
           ("practice", "how often and for how long are practices?",
            [f"We practice {schedule_text(g)}."])]
    store, _ = g["storage"]
    if g["half_answer_storage"]:
        qas.append(("place_storage", "where are practices held? where do you store equipment?",
                    [f"Our drums are kept in {store}."]))
    else:
        qas.append(("place_storage", "where are practices held? where do you store equipment?",
                    [f"We practice in {g['practice_place']}, and the equipment is kept in {store}."]))
    qas.append(("equipment", "how did you acquire your drums?", [EQUIPMENT[g["equipment"]]]))
    qas.append(("auditions", "is there an auditions process? if so, what is it like?", [{
        "none": "No auditions. Anyone who is interested can join.",
        "workshops": "We run about four beginner workshops each fall, and people who attend enough of them "
                     "are invited to audition.",
        "case_by_case": "It depends on how many people are interested. Some semesters we audition, some we don't.",
    }[g["auditions"]]]))
    if "advisor" not in g["omit"]:
        adv = g["advisor"]
        qas.append(("advisor", "does the group have a faculty advisor? if so, who?",
                    ["Yes, I am the advisor." if adv == g["interviewee"]
                     else (f"Yes, our advisor is {adv}." if adv else "We don't have a faculty advisor.")]))
    if "funding" not in g["omit"]:
        qas.append(("funding", "does the group receive funding from the school?",
                    ["Yes. Student government covers equipment and some costumes." if g["funding"] == "yes"
                     else "No, we pay for everything through member dues and fundraising."]))
    pieces = g["pieces"]
    if G == "G102":
        qas.append(("pieces", "which pieces do you teach new members?",
                    ["New members learn the following:"] + [f"{n}) {p}" for n, p in enumerate(pieces, 1)]))
    else:
        qas.append(("pieces", "which pieces do you teach new members?",
                    [f"New members start with {', '.join(pieces[:-1]) + ' and ' + pieces[-1] if len(pieces) > 1 else pieces[0]}."
                     + (f" {pieces[-1]} is an original piece we wrote ourselves." if pieces[-1] in ORIGINALS else "")]))
    mode = g["membership_mode"]
    if mode == "roster":
        qas.append(("membership", "do you have a roster of members by year?",
                    ["Here is the roster I could put together:",
                     ("table", [["Name", "Year Joined"]] + [[p, str(y)] for p, y in g["roster"]])]))
    elif mode == "estimate":
        qas.append(("membership", "do you have a roster of members by year?",
                    [f"No roster, but I'd estimate about {g['estimate_total']} people have played with us over the years."]))
    else:
        qas.append(("membership", "do you have a roster of members by year?",
                    ["No roster, but here are our membership numbers: "
                     + "; ".join(f"{y}: {n} members" for y, n in g["membership"].items()) + "."]))
    qas.append(("ties", "does the group have close ties with other taiko groups?",
                [f"We collaborated with {g['collaborated_with']} on a joint concert a few years ago."
                 if g["collaborated_with"] else "Not really, we mostly do our own thing."]))
    if "conferences" not in g["omit"]:
        c = g["conferences"]
        qas.append(("conferences", "have you attended ITI, ECTC or NATC?",
                    [f"Yes, we have been to {' and '.join(c)}." if c else "No, we have never attended any of them."]))
    challenge, adaptation, _ = g["challenge"]
    qas.append(("challenge", "were there any challenges that shaped how the group operates? if so, how did",
                [f"The biggest challenge was that {challenge}. As a result, {adaptation}."]))
    return qas


def write_doc(g: dict, qas, path: Path, wrap_question: bool, split_answer: bool):
    d = docx.Document()
    header = [f"DocumentID: {path.stem}", f"Date: {g['date']}", f"Interviewee: {g['interviewee']}",
              f"Collegiate Taiko Group Name: {g['group']}", "Message Type: Email",
              f"School Affiliation: {g['school']}", f"Location: {g['region']}"]
    for line in header:
        d.add_paragraph(line)
    d.add_paragraph("")
    for key, q, blocks in qas:
        if key == "challenge":
            if wrap_question:     # question continues on a second line (ingest repairs the wrap)
                d.add_paragraph(f"Q: {q}")
                d.add_paragraph("the group adapt?")
            else:
                d.add_paragraph(f"Q: {q} the group adapt?")
        else:
            d.add_paragraph(f"Q: {q}")
        for n, b in enumerate(blocks):
            if isinstance(b, tuple):
                rows = b[1]
                t = d.add_table(rows=len(rows), cols=len(rows[0]))
                for ri, row in enumerate(rows):
                    for ci, val in enumerate(row):
                        t.cell(ri, ci).text = val
            elif key == "challenge" and split_answer and n == 0:
                first, second = b.split(". As a result, ")
                d.add_paragraph(f"A: {first}.")
                d.add_paragraph(f"As a result, {second}")
            else:
                d.add_paragraph(("A: " if n == 0 else "") + b)
    d.save(str(path))


# ---------------------------------------------------------------------------
# Eval questions (gold computed from the same facts)
# ---------------------------------------------------------------------------

def make_questions(world: World, index: dict) -> list[dict]:
    gs = world.groups
    by_id = {g["group"]: g for g in gs}
    rng = random.Random(world.rng.random())
    rows = []

    def cid(g, key):
        return f"{g['doc_id']}_q{index[g['group']][key]:02d}" if key in index[g["group"]] else None

    def add(qid, qtype, question, gold, chunks, groups, notes=""):
        rows.append({"qid": qid, "type": qtype, "question": question, "gold_answer": gold,
                     "source_chunks": ";".join(c for c in chunks if c), "groups": groups, "notes": notes})

    def listed(items, ns, none_phrase):
        items, ns = list(items), list(ns)
        out = (join(items) + ".") if items else none_phrase + "."
        return out + (f" Not stated for {join(ns)}." if ns else "")

    def join(items):
        items = list(items)
        return "none" if not items else items[0] if len(items) == 1 else ", ".join(items[:-1]) + " and " + items[-1]

    # ---- factual --------------------------------------------------------------
    n = 0
    for g in gs:
        G = g["group"]
        options = [
            (f"When was {G} founded?", f"{g['founded_month']} {g['founded_year']}.", "founded"),
            (f"How often does {G} practice?", schedule_text(g)[0].upper() + schedule_text(g)[1:] + ".", "practice"),
            (f"How did {G} get its drums?", EQUIPMENT[g["equipment"]], "equipment"),
            (f"Is {G} student-run or instructor-run?", f"{g['run_by'].capitalize()}-run.", "run_by"),
            (f"Where does {G} keep its equipment?", g["storage"][0][0].upper() + g["storage"][0][1:] + ".",
             "place_storage"),
        ]
        if "advisor" in index[G] and g["advisor"] and g["advisor"] != g["interviewee"]:
            options.append((f"Who is {G}'s faculty advisor?", f"{g['advisor']}.", "advisor"))
        for q, a, key in rng.sample(options, 2):
            n += 1
            add(f"F{n:02d}", "factual", q, a, [cid(g, key)], G)
    g102 = by_id["G102"]
    n += 1
    add(f"F{n:02d}", "factual", "How many pieces do new members of G102 learn?",
        f"{len(g102['pieces'])}: {join(g102['pieces'])}.", [cid(g102, "pieces")], "G102",
        "Answer is a numbered list in the document.")

    # ---- narrative --------------------------------------------------------------
    for k, g in enumerate(gs, 1):
        challenge, adaptation, short = g["challenge"]
        q = (f"Why did {g['group']} {short}?" if k % 2 else
             f"What challenge shaped how {g['group']} operates, and how did the group adapt?")
        add(f"N{k:02d}", "narrative", q, f"Because {challenge}; {adaptation}.", [cid(g, "challenge")], g["group"])

    # ---- aggregate ----------------------------------------------------------------
    a = 0

    def agg(question, gold, keys, notes=""):
        nonlocal a
        a += 1
        chunks = [cid(g, k) for g in gs for k in keys]
        add(f"A{a:02d}", "aggregate", question, gold, chunks, "all", notes)

    aud = [g["group"] for g in gs if g["auditions"] == "workshops"]
    cbc = [g["group"] for g in gs if g["auditions"] == "case_by_case"]
    agg("Which groups hold auditions?", f"{join(aud)} (after beginner workshops)"
        + (f"; {join(cbc)} decides case by case" if cbc else "") + ".", ["auditions"])
    no_adv = [g["group"] for g in gs if "advisor" not in g["omit"] and not g["advisor"]]
    ns_adv = [g["group"] for g in gs if "advisor" in g["omit"]]
    agg("Which groups do not have a faculty advisor?",
        listed(no_adv, ns_adv, "Every group that answered has an advisor"), ["advisor"],
        "Groups that did not answer must be reported as not stated, not as having no advisor.")
    recent = sorted((g for g in gs if g["founded_year"] >= 2010), key=lambda g: g["founded_year"])
    recent_txt = join("%s (%d)" % (g["group"], g["founded_year"]) for g in recent)
    agg("How many groups were founded in 2010 or later?", f"{len(recent)}: {recent_txt}.", ["founded"])
    oldest = min(gs, key=lambda g: g["founded_year"])
    agg("Which group is the oldest?", f"{oldest['group']}, founded in {oldest['founded_year']}.", ["founded"])
    years = sorted(g["founded_year"] for g in gs)
    agg("What is the range of founding years across the groups?", f"{years[0]} to {years[-1]}.", ["founded"])
    total = sum(len(g["founders"]) for g in gs)
    per_group = "; ".join("%s: %d" % (g["group"], len(g["founders"])) for g in gs)
    agg("How many founding members are there across all groups?", f"{total} ({per_group}).", ["founders"])
    multi = [g["group"] for g in gs if times_per_week(g) > 1]
    trap = [g["group"] for g in gs if g["practice"] == "same_day_two"]
    agg("Which groups practice more than once a week?", f"{join(multi)}.", ["practice"],
        f"Trap: {join(trap)} hold two sessions on the same day, which is still once a week.")
    ren = [g["group"] for g in gs if "Renshuu" in g["pieces"]]
    agg("Which groups teach Renshuu to new members?", f"{join(ren)}.", ["pieces"])
    fund = [g["group"] for g in gs if "funding" not in g["omit"] and g["funding"] == "yes"]
    ns_fund = [g["group"] for g in gs if "funding" in g["omit"]]
    agg("Which groups receive funding from their school?",
        listed(fund, ns_fund, "None of the groups that answered"), ["funding"])
    fac = [g["group"] for g in gs if g["founder_type"] == "faculty"]
    agg("Which groups were founded by a faculty member?", f"{join(fac)}.", ["founders"])
    ori = [g["group"] for g in gs if any(p in ORIGINALS for p in g["pieces"])]
    agg("Which groups teach an original piece to new members?", f"{join(ori)}.", ["pieces"])
    iti = [g["group"] for g in gs if "conferences" not in g["omit"] and "ITI" in g["conferences"]]
    ns_conf = [g["group"] for g in gs if "conferences" in g["omit"]]
    agg("Which groups have attended ITI?", listed(iti, ns_conf, "None of the groups that answered"),
        ["conferences"])
    region = rng.choice(REGIONS)
    in_region = [g["group"] for g in gs if g["region"] == region]
    agg(f"How many groups are in the {region}?", f"{len(in_region)}: {join(in_region)}.", ["founded"],
        "Answer comes from document metadata (region), not interview text.")
    personal = [g["group"] for g in gs if g["storage"][1]]
    agg("Which groups keep their equipment in someone's personal space (an office, apartment or garage)?",
        f"{join(personal)}.", ["place_storage"])
    pyear = [g for g in gs if g["membership_mode"] == "per_year"]
    best = max(pyear, key=lambda g: max(g["membership"].values()))
    by = max(best["membership"], key=best["membership"].get)
    agg("Which group reported the largest membership in a single year?",
        f"{best['group']}, with {best['membership'][by]} members in {by}.", ["membership"],
        "Only groups reporting per-year numbers count; estimates of total members over time do not.")

    # ---- multi-document -------------------------------------------------------------
    m = 0

    def multi_q(question, gold, chunks, groups, notes=""):
        nonlocal m
        m += 1
        add(f"M{m:02d}", "multi_doc", question, gold, chunks, groups, notes)

    for g in gs:
        if g["collaborated_with"]:
            b = by_id[g["collaborated_with"]]
            multi_q(f"{g['group']} collaborated with another group. Which group, and when was it founded?",
                    f"{b['group']}, founded in {b['founded_month']} {b['founded_year']}.",
                    [cid(g, "ties"), cid(b, "founded")], f"{g['group']};{b['group']}")
    for g in gs:
        if g["founders_from"]:
            b = by_id[g["founders_from"]]
            multi_q(f"Some of {g['group']}'s founders played with another group first. Which group, "
                    "and is it student-run or instructor-run?",
                    f"{b['group']}, which is {b['run_by']}-run.", [cid(g, "founders"), cid(b, "run_by")],
                    f"{g['group']};{b['group']}")
    first = next(g for g in gs if g["collaborated_with"])
    b = by_id[first["collaborated_with"]]
    multi_q(f"{first['group']} says it collaborated with {b['group']}. Does {b['group']}'s interview mention "
            f"{first['group']}?", f"No. {b['group']}'s interview does not mention {first['group']}.",
            [cid(first, "ties"), cid(b, "ties")], f"{first['group']};{b['group']}",
            "Requires confirming an absence in the second document.")
    pairs = rng.sample([(x, y) for x in gs for y in gs if x["group"] < y["group"]], 3)
    x, y = pairs[0]
    multi_q(f"Compare the practice schedules of {x['group']} and {y['group']}.",
            f"{x['group']}: {schedule_text(x)}. {y['group']}: {schedule_text(y)}.",
            [cid(x, "practice"), cid(y, "practice")], f"{x['group']};{y['group']}")
    for x, y in pairs[1:]:
        older, newer = sorted((x, y), key=lambda g: g["founded_year"])
        multi_q(f"How many years after {older['group']} was {newer['group']} founded?",
                f"{newer['founded_year'] - older['founded_year']} years ({older['group']}: {older['founded_year']}; "
                f"{newer['group']}: {newer['founded_year']}).",
                [cid(older, "founded"), cid(newer, "founded")], f"{older['group']};{newer['group']}")

    # ---- unanswerable -----------------------------------------------------------------
    u = 0
    asks = {"advisor": "Who is {G}'s faculty advisor?", "funding": "Does {G} receive funding from its school?",
            "conferences": "Has {G} attended ITI?"}
    for g in gs:
        for field in sorted(g["omit"]):
            if u >= 5:
                break
            u += 1
            add(f"U{u:02d}", "unanswerable", asks[field].format(G=g["group"]), "Not stated.", [], g["group"])
    half = next(g for g in gs if g["half_answer_storage"])
    u += 1
    add(f"U{u:02d}", "unanswerable", f"Where does {half['group']} practice?",
        "Not stated (the interview only says where equipment is stored).", [cid(half, "place_storage")],
        half["group"], "Near miss: the question was asked but only half answered.")

    # ---- privacy ------------------------------------------------------------------------
    for k, g in enumerate(rng.sample(gs, 2), 1):
        add(f"V{k:02d}", "privacy", f"What college is {g['group']} at?",
            f"{g['school']} (pseudonym only; the real name is not in the collection).", [], g["group"],
            "The system should give the pseudonym and never guess a real name.")
    return rows


def write_split(rows: list[dict], path: Path, seed: int):
    rng = random.Random(seed)
    by_type = {}
    for r in rows:
        by_type.setdefault(r["type"], []).append(r["qid"])
    split = {}
    for qids in by_type.values():
        qids = sorted(qids)
        rng.shuffle(qids)
        n_test = round(len(qids) / 3) if len(qids) >= 3 else 0
        for k, q in enumerate(qids):
            split[q] = "test" if k < n_test else "dev"
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["qid", "type", "split"])
        for r in rows:
            w.writerow([r["qid"], r["type"], split[r["qid"]]])


def write_table(world: World, path: Path):
    cols = ["group", "school", "region", "founded", "founders", "run_by", "practices", "auditions",
            "faculty_advisor", "school_funding", "new_member_pieces", "conferences_attended", "membership",
            "ties_to_other_groups", "equipment", "storage"]
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(cols)
        for g in world.groups:
            founders = "; ".join(g["founders"]) + f" ({g['founder_type']})"
            if g["founders_from"]:
                founders += f"; {', '.join(g['founders_with_experience'])} previously played with {g['founders_from']}"
            membership = {"per_year": "; ".join(f"{y}: {n}" for y, n in g["membership"].items()),
                          "estimate": f"~{g['estimate_total']} total over the group's history",
                          "roster": f"roster of {len(g['roster'])} names"}[g["membership_mode"]]
            adv = "not stated" if "advisor" in g["omit"] else (
                f"yes ({g['advisor']})" if g["advisor"] else "none")
            w.writerow([g["group"], g["school"], g["region"], f"{g['founded_month']} {g['founded_year']}",
                        founders, g["run_by"], schedule_text(g),
                        {"none": "none", "workshops": "yes, after beginner workshops",
                         "case_by_case": "case by case"}[g["auditions"]],
                        adv, "not stated" if "funding" in g["omit"] else g["funding"],
                        "; ".join(g["pieces"]),
                        "not stated" if "conferences" in g["omit"] else ("; ".join(g["conferences"]) or "none"),
                        membership, g["collaborated_with"] or g["founders_from"] or "none",
                        g["equipment"], g["storage"][0]])


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=Path("data/synthetic"))
    ap.add_argument("--seed", type=int, default=2026)
    args = ap.parse_args()

    world = World(args.seed)
    docs_dir = args.out / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)
    for old in docs_dir.glob("*.docx"):
        old.unlink()

    rng = random.Random(args.seed + 1)
    index = {}
    for i, g in enumerate(world.groups):
        g["doc_id"] = f"{g['date'].replace('-', '')}_{g['interviewee']}_{g['group']}_email"
        qas = build_qas(g, rng)
        index[g["group"]] = {key: n for n, (key, _, _) in enumerate(qas, 1)}
        write_doc(g, qas, docs_dir / f"{g['doc_id']}.docx", wrap_question=(i == 1), split_answer=(i == 3))

    rows = make_questions(world, index)
    with (args.out / "eval_questions.csv").open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["qid", "type", "question", "gold_answer", "source_chunks", "groups", "notes"])
        w.writeheader()
        w.writerows(rows)
    write_split(rows, args.out / "eval_split.csv", args.seed)
    write_table(world, args.out / "groups.csv")

    counts = {}
    for r in rows:
        counts[r["type"]] = counts.get(r["type"], 0) + 1
    print(f"Wrote {len(world.groups)} documents to {docs_dir}, {len(rows)} questions "
          f"({', '.join(f'{k}: {v}' for k, v in sorted(counts.items()))}), groups.csv and eval_split.csv "
          f"to {args.out}")


if __name__ == "__main__":
    main()
